from __future__ import annotations

import copy
import hashlib
import json
import shutil
import sqlite3
import subprocess
import sys
import tempfile
import unittest
import importlib.util
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
from qce_compat.database import LocalDatabase
from qce_compat.media import MediaExporter, MediaLibrary
from qce_compat.output import document, write_excel
from qce_compat.protocol import Decoder, fields, identifier, resolve_quotes
from tests.fixtures import LARGE_ID, START, body, cached, create, element, field, text


class ProtocolTests(unittest.TestCase):
    def test_truncated_protobuf_is_explicitly_preserved(self):
        decoder = Decoder()
        decoded = decoder.body(b"\x12\x08abc")
        self.assertEqual(decoded[0]["type"], "unknown")
        self.assertEqual(decoded[0]["data"]["rawHex"], "1208616263")
        self.assertTrue(decoder.warnings)

    def test_large_id_remains_exact_string(self):
        self.assertEqual(identifier(LARGE_ID), "9007199254740993123")
        self.assertEqual(identifier(-1), "18446744073709551615")

    def test_reply_uses_correct_time_and_sender_fields(self):
        result = Decoder().body(body(element(7, **{"47401": LARGE_ID, "47403": 123456789, "47404": START})))
        self.assertEqual(result[0]["data"]["senderUin"], "123456789")
        self.assertEqual(result[0]["data"]["timestamp"], START)

    def test_nested_cached_messages_keep_all_media(self):
        decoder = Decoder()
        inner = cached(2, text("nested"), element(5, **{"45402": "movie.mp4"}))
        outer = cached(1, element(16), nested=[inner])
        messages = decoder.cached(field(40900, outer), "scope")
        nested = messages[0]["content"]["elements"][0]["data"]["messages"]
        self.assertEqual([item["type"] for item in nested[0]["content"]["elements"]], ["text", "video"])
        self.assertTrue(messages[0]["id"].startswith("forward:scope/0:"))


class FixtureTests(unittest.TestCase):
    def setUp(self):
        (ROOT / "test-output").mkdir(exist_ok=True)
        self.temporary = tempfile.TemporaryDirectory(prefix="qce-test-", dir=ROOT / "test-output")
        self.root = Path(self.temporary.name)
        self.fixture = create(self.root)
        self.db = LocalDatabase(self.fixture["source"])
        self.session = next(session for session in self.db.sessions(["111111"]))

    def tearDown(self):
        self.db.close()
        self.temporary.cleanup()

    def messages(self):
        result = self.db.messages(self.session)
        resolve_quotes(result, lambda quote: self.db.lookup_quote(quote, self.session))
        return result

    def test_scope_and_date_filter(self):
        messages = self.db.messages(self.session, since=START + 1, until=START + 3)
        self.assertEqual([item["id"] for item in messages], ["1002", "1003"])
        self.assertEqual(self.session["name"], "演示群:一")

    def test_reply_to_out_of_range_message_gets_media_without_fake_jump(self):
        messages = self.db.messages(self.session, since=START + 1, until=START + 2)
        resolve_quotes(messages, lambda quote: self.db.lookup_quote(quote, self.session))
        quote = messages[0]["content"]["elements"][0]["data"]
        self.assertFalse(quote["sourceAvailable"])
        self.assertEqual(quote["referencedMessageId"], "")
        self.assertEqual(quote["reference"]["messageId"], str(LARGE_ID))
        self.assertIn("image", [item["type"] for item in quote["previewElements"]])

    def test_timeline_quote_jumps_to_real_id_not_snapshot_id(self):
        quote = self.messages()[1]["content"]["elements"][0]["data"]
        self.assertTrue(quote["sourceAvailable"])
        self.assertEqual(quote["referencedMessageId"], str(LARGE_ID))
        self.assertNotEqual(quote["referencedMessageId"], quote["reference"]["snapshotId"])

    def test_resources_deduplicate_and_failures_keep_tree(self):
        messages = self.messages()
        destination = self.root / "converted"
        report = MediaExporter(MediaLibrary([self.fixture["media"]]), destination).process(messages)
        self.assertEqual(report["saved"], 4)
        self.assertEqual(report["failed"], 1)
        self.assertEqual(len(report["failures"][0]["occurrences"]), 2)
        self.assertEqual(len(report["failures"][0]["occurrences"][0]["tree"]), 2)
        files = [path for path in destination.rglob("*") if path.is_file()]
        self.assertEqual(len(files), 4)
        snapshot = messages[4]["content"]["elements"][0]["data"]["previewElements"][0]
        self.assertEqual(snapshot["type"], "video")
        self.assertTrue(snapshot["localPath"].startswith("videos/"))
        self.assertNotIn("sourcePaths", json.dumps(messages))

    def test_ambiguous_same_basename_is_not_silently_chosen(self):
        first, second = self.fixture["media"] / "a", self.fixture["media"] / "b"
        first.mkdir(); second.mkdir()
        (first / "same.bin").write_bytes(b"first file")
        (second / "same.bin").write_bytes(b"another file")
        self.assertIsNone(MediaLibrary([self.fixture["media"]]).find("file", {"filename": "same.bin"}))

    def test_paths_outside_media_roots_are_never_copied(self):
        secret = self.root / "outside.txt"
        secret.write_text("outside", encoding="utf-8")
        library = MediaLibrary([self.fixture["media"]])
        self.assertIsNone(library.find("file", {"sourcePaths": [str(secret)]}))

    def test_qq_crc64_cache_without_original_filename(self):
        # Cache name calculated from the QQNT_Export reference algorithm for this PNG.
        path = self.fixture["media"] / "chatpic" / "chatraw" / "841" / "Cache_-33cba5e1fa0b7841"
        path.parent.mkdir(parents=True)
        path.write_bytes(self.fixture["image"].read_bytes())
        self.fixture["image"].unlink()
        found = MediaLibrary([self.fixture["media"]]).find("image", {"md5": "5af9df7cae4e3fba8cca7c631a4a178d"})
        self.assertIsNotNone(found)
        self.assertEqual(found[0], path)

    def test_incorrect_hash_does_not_save_a_thumbnail_as_original(self):
        library = MediaLibrary([self.fixture["media"]])
        self.assertIsNone(library.find("image", {"md5": "aa" * 16, "sourcePaths": [str(self.fixture["image"])]}))

    def test_original_database_is_read_only(self):
        before = hashlib.sha256(self.fixture["source"].read_bytes()).hexdigest()
        self.messages()
        with self.assertRaises(sqlite3.OperationalError):
            self.db.connection.execute("DELETE FROM group_msg_table")
        self.assertEqual(before, hashlib.sha256(self.fixture["source"].read_bytes()).hexdigest())

    def test_minimal_schema_does_not_invent_missing_columns(self):
        source = self.root / "minimal.db"
        with sqlite3.connect(source) as connection:
            connection.execute('CREATE TABLE c2c_msg_table ("40001", "40021", "40800")')
            connection.execute('INSERT INTO c2c_msg_table VALUES (?,?,?)', (1, "u_min", body(text("minimal"))))
        connection.close()
        database = LocalDatabase(source)
        try:
            session = next(database.sessions())
            self.assertEqual(database.messages(session)[0]["content"]["text"], "minimal")
        finally:
            database.close()

    def run_cli(self, *arguments, expected=0):
        command = [sys.executable, str(ROOT / "scripts/ntqq_export_qce.py"), "--src", str(self.fixture["source"]),
                   "--media-root", str(self.fixture["media"]), "--outdir", str(self.root / "exported"),
                   "--config", str(self.root / "no-config.json"), *arguments]
        result = subprocess.run(command, capture_output=True, text=True, encoding="utf-8", errors="replace")
        self.assertEqual(result.returncode, expected, result.stdout + result.stderr)
        return result

    def test_json_text_excel_html_zip_end_to_end(self):
        self.run_cli("--only", "111111", "--zip")
        zips = list((self.root / "exported").rglob("*.zip"))
        self.assertEqual(len(zips), 1)
        with zipfile.ZipFile(zips[0]) as archive:
            names = archive.namelist()
            self.assertEqual(len(names), len({name.casefold() for name in names}))
            self.assertFalse(any("partial" in name for name in names))
            report = json.loads(archive.read("resources/media-report.json"))
            self.assertEqual(report["failed"], 1)
            json_name = next(name for name in names if name.endswith(".json") and not name.startswith("resources/"))
            data = json.loads(archive.read(json_name))
            self.assertEqual(data["chatInfo"]["peerUid"], "111111")
            self.assertEqual(data["messages"][0]["id"], str(LARGE_ID))
            self.assertEqual(data["statistics"]["totalMessages"], 8)
            html = archive.read(next(name for name in names if name.endswith(".html"))).decode()
            self.assertIn('class="forward-card forward-card-expandable', html)
            self.assertIn('reply-content-video', html)
            self.assertTrue('data-reply-to="msg-' + str(LARGE_ID) in html, "引用应跳转到真实消息 ID")
            self.assertNotIn('src="https://unpkg.com', html)
            self.assertIn("[2026-10-19]", zips[0].name)
            self.assertTrue(any(name.startswith("resources/videos/") and name.endswith(".mp4") for name in names))
            self.assertTrue(any(name.endswith(".xlsx") for name in names))
        self.assertFalse(list((self.root / "exported").rglob(".qce-stage-*")))

    def test_each_session_gets_its_own_package(self):
        self.run_cli("--formats", "json,txt", "--zip")
        self.assertEqual(len(list((self.root / "exported").rglob("*.zip"))), 3)
        self.assertEqual(len({path.parent for path in (self.root / "exported").rglob("*.zip")}), 3)

    def test_repeated_export_does_not_overwrite_previous_package(self):
        self.run_cli("--only", "111111", "--formats", "json", "--zip")
        first = next((self.root / "exported").rglob("*.zip"))
        checksum = hashlib.sha256(first.read_bytes()).hexdigest()
        self.run_cli("--only", "111111", "--formats", "json", "--zip")
        self.assertEqual(len(list((self.root / "exported").rglob("*.zip"))), 2)
        self.assertEqual(hashlib.sha256(first.read_bytes()).hexdigest(), checksum)

    def test_excel_literal_formulas_and_long_messages(self):
        from openpyxl import load_workbook
        messages = self.messages()[:1]
        messages[0]["sender"]["name"] = "=1+1"
        messages[0]["content"]["elements"] = [{"type": "text", "data": {"text": "长文本" * 22000}}]
        report = {"saved": 0, "failed": 0, "total": 0, "resources": [], "failures": []}
        data = document(messages, {"name": "demo", "type": "group"}, report)
        path = self.root / "long.xlsx"
        write_excel(path, data)
        workbook = load_workbook(path, read_only=True)
        try:
            sheet = workbook["聊天记录"]
            rows = list(sheet.rows)
            self.assertEqual(rows[1][2].value, "=1+1")
            self.assertEqual(rows[1][2].data_type, "s")
            self.assertEqual(len(rows), 4)
            self.assertEqual(workbook.sheetnames, ["聊天记录", "统计信息", "发送者统计", "资源列表"])
        finally:
            workbook.close()

    def test_encrypted_input_is_rejected_without_output(self):
        encrypted = self.root / "encrypted.db"
        encrypted.write_bytes(b"not a plaintext SQLite database")
        self.run_cli("--src", str(encrypted), "--formats", "json", expected=1)
        self.assertFalse((self.root / "exported").exists())

    @unittest.skipUnless(importlib.util.find_spec("sqlcipher3"), "SQLCipher dependency is optional for conversion-only tests")
    def test_upstream_sqlcipher_can_decrypt_a_synthetic_qqnt_database(self):
        from sqlcipher3 import dbapi2 as cipher
        from ntqq_decrypt import backup_db, strip_header, probe_key, export_plain
        raw_key, salt = "01" * 32, "02" * 16  # Public synthetic test key, never an account key.
        encrypted = self.root / "cipher.db"
        connection = cipher.connect(self.fixture["source"])
        try:
            connection.execute("ATTACH DATABASE ? AS encrypted KEY ?", (str(encrypted), f"x'{raw_key}{salt}'"))
            connection.execute("PRAGMA encrypted.cipher_page_size=4096")
            connection.execute("SELECT sqlcipher_export('encrypted')")
            connection.execute("DETACH DATABASE encrypted")
        finally:
            connection.close()
        wrapped = self.root / "qq-wrapped.db"
        wrapped.write_bytes(b"\0" * 1024 + encrypted.read_bytes())
        work = self.root / "crypto-work"
        work.mkdir()
        backed_up = backup_db(str(wrapped), str(work))
        stripped = strip_header(backed_up, str(work / "clear.db"))
        variant, opened, errors = probe_key(stripped, raw_key, salt)
        self.assertIsNotNone(opened, "Synthetic SQLCipher round trip failed")
        try:
            target = self.root / "decrypted.db"
            export_plain(opened, str(target))
        finally:
            opened.close()
        reader = LocalDatabase(target, aux_dir=self.root)
        try:
            self.assertEqual(len(list(reader.sessions())), 3)
        finally:
            reader.close()


if __name__ == "__main__":
    (ROOT / "test-output").mkdir(exist_ok=True)
    unittest.main()
