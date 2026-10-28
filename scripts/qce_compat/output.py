"""QCE JSON, text, workbook and ZIP output, sharing a single resource tree."""
from __future__ import annotations

import collections
import json
import math
import re
import zipfile
from datetime import datetime, timezone
from pathlib import Path

from . import VERSION
from .media import iter_media

TYPE_LABELS = {"text": "文本消息", "image": "图片消息", "video": "视频消息", "audio": "语音消息", "file": "文件消息",
               "mixed": "混合消息", "forward": "合并转发", "system": "系统消息", "json": "卡片消息", "face": "表情消息"}


def document(messages, chat_info, report):
    kinds, senders, resources = collections.Counter(), {}, collections.Counter()
    for message in messages:
        kinds[message["type"]] += 1
        sender = message["sender"]
        key = sender.get("uid") or sender.get("uin") or sender["name"]
        entry = senders.setdefault(key, {"uid": sender.get("uid", ""), "uin": sender.get("uin", ""), "name": sender["name"], "messageCount": 0})
        entry["messageCount"] += 1
    for entry in senders.values():
        entry["percentage"] = round(entry["messageCount"] * 100 / len(messages), 2) if messages else 0
    for resource in report["resources"]:
        resources[resource["type"]] += 1
    times = sorted(message["timestamp"] for message in messages if message["timestamp"] > 0)
    def iso(value):
        return datetime.fromtimestamp(value / 1000, timezone.utc).isoformat(timespec="milliseconds").replace("+00:00", "Z") if value else ""
    stats = {"totalMessages": len(messages), "timeRange": {"start": iso(times[0]) if times else "", "end": iso(times[-1]) if times else "",
             "durationDays": math.ceil((times[-1] - times[0]) / 86_400_000) if times else 0},
             "messageTypes": dict(kinds), "senders": sorted(senders.values(), key=lambda value: -value["messageCount"]),
             "resources": {"total": report["total"], "byType": dict(resources), "totalSize": sum(value["size"] for value in report["resources"])}}
    return {"metadata": {"name": "qq-nt-export / QQ Chat Exporter compatible", "version": VERSION,
                         "copyright": "q957506633; QQ Chat Exporter contributors", "exportedAt": datetime.now(timezone.utc).isoformat()},
            "chatInfo": chat_info, "statistics": stats, "messages": messages,
            "exportOptions": {"includedFields": ["sender", "content", "resources", "reply", "forward"], "filters": {},
                              "options": {"dataSource": "local-sqlite", "mediaSource": "local-cache", "schema": "qce-clean-message"}},
            "mediaReport": {key: value for key, value in report.items() if key != "resources"}}


def write_json(path, value):
    with Path(path).open("w", encoding="utf-8", newline="\n") as stream:
        json.dump(value, stream, ensure_ascii=False, indent=2)
        stream.write("\n")


def describe_elements(elements, indent=""):
    lines = []
    for element in elements:
        kind, data = element["type"], element["data"]
        if kind == "reply":
            lines.append(f"{indent}回复: {data.get('senderName', '')} - {data.get('content', '')}")
            for preview in data.get("previewElements", []):
                if preview.get("type") in ("image", "video", "audio", "file"):
                    label = preview.get("fileName") or preview.get("filename") or preview.get("text", "")
                    state = " [媒体导出失败]" if preview.get("exportState") else ""
                    link = " (resources/" + preview["localPath"] + ")" if preview.get("localPath") else ""
                    lines.append(f"{indent}  引用资源: {label}{link}{state}")
        elif kind == "forward":
            lines.append(f"{indent}[聊天记录] {data.get('title', '')}")
            if not data.get("messages"):
                lines.append(indent + "  [本地未缓存该合并记录的完整内容]")
            for message in data.get("messages", []):
                lines.append(f"{indent}  {message.get('time', '')} {message['sender']['name']}:")
                lines.extend(describe_elements(message["content"]["elements"], indent + "    "))
        elif kind in ("image", "video", "audio", "file"):
            state = " [媒体导出失败]" if data.get("exportState") else ""
            link = " (resources/" + data["localPath"] + ")" if data.get("localPath") else ""
            lines.append(f"{indent}[{TYPE_LABELS.get(kind, kind)}] {data.get('filename', '')}{link}{state}")
        else:
            text = data.get("text") or data.get("name") or data.get("title") or f"[{kind}]"
            lines.append(indent + str(text))
    return lines


def write_text(path, data):
    chat, messages, stats = data["chatInfo"], data["messages"], data["statistics"]
    with Path(path).open("w", encoding="utf-8-sig", newline="\n") as stream:
        def line(value=""):
            stream.write(value + "\n")
        line("[QQChatExporter 兼容导出 / qq-nt-export]")
        line("===============================================")
        line("           QQ聊天记录导出文件")
        line("===============================================")
        line(f"聊天名称: {chat['name']}")
        line(f"聊天类型: {'群聊' if chat['type'] == 'group' else '私聊'}")
        line(f"消息总数: {len(messages)}")
        line(f"时间范围: {stats['timeRange']['start']} 至 {stats['timeRange']['end']}")
        line()
        for index, message in enumerate(messages, 1):
            line(f"[{index}]")
            line(message["sender"]["name"] + ":")
            line("时间: " + message["time"])
            for text in describe_elements(message["content"]["elements"]):
                line(text)
            line()
        line("===============================================")
        line(f"总计导出 {len(messages)} 条消息；媒体已保存 {data['mediaReport']['saved']}，失败 {data['mediaReport']['failed']}")


def write_excel(path, data):
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    workbook = Workbook(write_only=True)
    illegal_xml = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")

    def sheet(name, headers, widths):
        worksheet = workbook.create_sheet(name)
        worksheet.freeze_panes = "A2"
        for index, width in enumerate(widths, 1):
            worksheet.column_dimensions[get_column_letter(index)].width = width
        cells = []
        for header in headers:
            cell = WriteOnlyCell(worksheet, header)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="4472C4")
            cells.append(cell)
        worksheet.append(cells)
        return worksheet

    def append(worksheet, values):
        cells = []
        for value in values:
            cell = WriteOnlyCell(worksheet, illegal_xml.sub("", value)[:32767] if isinstance(value, str) else value)
            if isinstance(value, str):
                cell.data_type = "s"  # Chat text beginning with '=' remains text.
            cells.append(cell)
        worksheet.append(cells)

    headers = ["序号", "时间", "发送者", "发送者QQ号", "消息类型", "消息内容", "是否撤回", "资源数量"]
    widths = [8, 28, 20, 18, 14, 80, 12, 12]
    messages_sheet = sheet("聊天记录", headers, widths)
    row_count, part = 1, 1
    for index, message in enumerate(data["messages"], 1):
        text = "\n".join(describe_elements(message["content"]["elements"]))
        chunks = [text[start:start + 32000] for start in range(0, len(text), 32000)] or [""]
        for chunk_index, chunk in enumerate(chunks, 1):
            if row_count >= 1_048_576:
                part += 1
                messages_sheet = sheet(f"聊天记录{part}", headers, widths)
                row_count = 1
            if len(chunks) > 1:
                chunk = f"（消息内容 {chunk_index}/{len(chunks)}）\n" + chunk
            append(messages_sheet, [index, message["time"], message["sender"]["name"], message["sender"].get("uin", ""),
                TYPE_LABELS.get(message["type"], message["type"]), chunk, "是" if message.get("recalled") else "否", len(message["content"]["resources"])])
            row_count += 1
    stats_sheet = sheet("统计信息", ["统计项目", "数值"], [24, 40])
    stats = data["statistics"]
    for item in [("消息总数", len(data["messages"])), ("开始时间", stats["timeRange"]["start"]), ("结束时间", stats["timeRange"]["end"]),
                 ("时间跨度(天)", stats["timeRange"]["durationDays"]), ("已保存媒体", data["mediaReport"]["saved"]), ("失败媒体", data["mediaReport"]["failed"])]:
        append(stats_sheet, item)
    for kind, count in stats["messageTypes"].items():
        append(stats_sheet, (TYPE_LABELS.get(kind, kind), count))
    sender_sheet = sheet("发送者统计", ["排名", "发送者", "QQ号", "UID", "消息数量", "占比(%)"], [8, 24, 18, 28, 14, 14])
    for index, sender in enumerate(stats["senders"], 1):
        append(sender_sheet, [index, sender["name"], sender.get("uin", ""), sender["uid"], sender["messageCount"], sender["percentage"]])
    resources_sheet = sheet("资源列表", ["序号", "时间", "发送者", "发送者QQ号", "资源类型", "文件名", "大小(字节)", "URL", "导出状态", "定位层级"], [8, 28, 24, 18, 14, 45, 18, 65, 18, 65])
    index = 0
    for message in data["messages"]:
        for owner, kind, resource, trail in iter_media(message):
            index += 1
            append(resources_sheet, [index, owner.get("time", ""), owner["sender"]["name"], owner["sender"].get("uin", ""), kind,
                resource.get("filename") or resource.get("fileName", ""), resource.get("size", 0),
                "resources/" + resource["localPath"] if resource.get("localPath") else "", "失败" if resource.get("exportState") else "已保存",
                " → ".join("引用" if branch["kind"] == "reply" else f"{branch['title']} / 第{branch['index']}条" for branch in trail)])
    workbook.save(path)


def archive(directory, destination):
    directory, destination = Path(directory).resolve(), Path(destination)
    seen = set()
    with zipfile.ZipFile(destination, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=6, allowZip64=True) as target:
        for path in sorted(directory.rglob("*")):
            if path.is_symlink():
                raise ValueError("导出目录含符号链接，已停止打包")
            relative = path.relative_to(directory).as_posix()
            key = relative.rstrip("/").casefold()
            if key in seen:
                raise ValueError("ZIP 出现重复路径：" + relative)
            seen.add(key)
            if path.is_file():
                target.write(path, relative)
            elif path.is_dir() and not any(path.iterdir()):
                target.writestr(relative + "/", b"")
